"""Export physical risk feature results for Client RE Fund to Excel.

Runs assess_physical_risk in both API (use_api_data=True) and static
(use_api_data=False) modes, then writes a five-sheet Excel workbook to
outputs/client_re_fund_physical_risk_feature_results.xlsx.

Usage:
    python -X utf8 sandbox/export_physical_risk_feature.py
"""

import datetime
import subprocess
import sys
from pathlib import Path

# ── Path setup ──────────────────────────────────────────────────────────────
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "backend"))

import openpyxl  # noqa: E402  (after path manipulation)
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.physical_risk import _region_type, assess_physical_risk  # noqa: E402
from app.data.sample_facilities import get_facilities_by_company  # noqa: E402

# ── Colour constants ─────────────────────────────────────────────────────────
_FILL_HEADER = PatternFill("solid", fgColor="BDD7EE")   # light blue
_FILL_HIGH   = PatternFill("solid", fgColor="FFCCCC")   # light red
_FILL_MEDIUM = PatternFill("solid", fgColor="FFFF99")   # light yellow
_FILL_LOW    = PatternFill("solid", fgColor="CCFFCC")   # light green

_FONT_BOLD = Font(bold=True)


def _autofit(ws: openpyxl.worksheet.worksheet.Worksheet) -> None:
    """Approximate column autofit by measuring cell content width."""
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 60)


def _header_row(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    row: int,
    values: list[str],
) -> None:
    for col, val in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=val)
        cell.font = _FONT_BOLD
        cell.fill = _FILL_HEADER


# ── Data collection ──────────────────────────────────────────────────────────
print("Client RE Fund 시설 데이터 로딩 중...")
facs = get_facilities_by_company("Client RE Fund")
print(f"  → {len(facs)}개 시설 확인")

print("API 모드 평가 실행 중 (use_api_data=True)...")
r_api = assess_physical_risk(
    "current_policies", 2030, use_api_data=True, facilities=facs
)
print("  → 완료")

print("정적 모드 평가 실행 중 (use_api_data=False)...")
r_static = assess_physical_risk(
    "current_policies", 2030, use_api_data=False, facilities=facs
)
print("  → 완료")

branch = subprocess.check_output(
    ["git", "branch", "--show-current"],
    text=True,
    cwd=str(Path(__file__).resolve().parent.parent),
).strip()
run_time = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")


# ── Derived aggregates ────────────────────────────────────────────────────────
def _sum_potential_loss(result: dict) -> float:
    total = 0.0
    for fac_r in result["facilities"]:
        for h in fac_r["hazards"]:
            total += h.get("potential_loss", 0)
    return total


def _sum_eal(result: dict) -> float:
    return sum(f["total_expected_annual_loss"] for f in result["facilities"])


total_pl_api    = _sum_potential_loss(r_api)
total_pl_static = _sum_potential_loss(r_static)
total_eal_api   = _sum_eal(r_api)
total_eal_static = _sum_eal(r_static)

eal_change_pct = (
    round((total_eal_api - total_eal_static) / total_eal_static * 100, 1)
    if total_eal_static != 0
    else "N/A"
)

api_warnings_str = "; ".join(
    sorted({w for f in r_api["facilities"] for w in f.get("api_warnings", [])})
) or "없음"


# ── Build workbook ────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()


# ── Sheet 1: Summary ──────────────────────────────────────────────────────────
ws1 = wb.active
ws1.title = "Summary"

summary_rows: list[tuple[str, object]] = [
    ("회사명",                   "Client RE Fund"),
    ("실행 시각",                run_time),
    ("브랜치명",                  branch),
    ("시나리오",                  "current_policies"),
    ("평가 연도",                 2030),
    ("use_api_data",             "True"),
    ("총 potential_loss (API)",  total_pl_api),
    ("총 potential_loss (Static)", total_pl_static),
    ("총 EAL (API)",             total_eal_api),
    ("총 EAL (Static)",          total_eal_static),
    ("EAL 변화 (%)",             eal_change_pct),
    ("api_warnings",             api_warnings_str),
    (
        "주요 변경 요약",
        "홍수·폭염·가뭄은 ERA5 좌표 기반으로 전환. 태풍·해수면 상승은 정적 권역 기반 유지.",
    ),
]

_numeric_labels = {
    "총 potential_loss (API)",
    "총 potential_loss (Static)",
    "총 EAL (API)",
    "총 EAL (Static)",
    "EAL 변화 (%)",
}

# Header row
ws1.cell(row=1, column=1, value="항목").font = _FONT_BOLD
ws1.cell(row=1, column=1).fill = _FILL_HEADER
ws1.cell(row=1, column=2, value="값").font = _FONT_BOLD
ws1.cell(row=1, column=2).fill = _FILL_HEADER

for row_idx, (label, value) in enumerate(summary_rows, start=2):
    label_cell = ws1.cell(row=row_idx, column=1, value=label)
    label_cell.font = _FONT_BOLD
    val_cell = ws1.cell(row=row_idx, column=2, value=value)
    if label in _numeric_labels and isinstance(value, (int, float)):
        val_cell.number_format = "#,##0.0" if label == "EAL 변화 (%)" else "#,##0"

ws1.column_dimensions["A"].width = 30
ws1.column_dimensions["B"].width = 80


# ── Sheet 2: Input_Facilities ─────────────────────────────────────────────────
ws2 = wb.create_sheet("Input_Facilities")

fac_cols = [
    "facility_id", "facility_name", "company", "sector", "location",
    "latitude", "longitude", "region_type",
    "assets_value", "annual_revenue", "ebitda",
    "current_emissions_scope1", "current_emissions_scope2", "current_emissions_scope3",
]
_numeric_fac_cols = {"assets_value", "annual_revenue", "ebitda"}

_header_row(ws2, 1, fac_cols)

for row_idx, fac in enumerate(facs, start=2):
    region_type = _region_type(fac["latitude"], fac["longitude"])
    row_values = {
        "facility_id":   fac["facility_id"],
        "facility_name": fac["name"],
        "company":       fac["company"],
        "sector":        fac["sector"],
        "location":      fac["location"],
        "latitude":      fac["latitude"],
        "longitude":     fac["longitude"],
        "region_type":   region_type,
        "assets_value":  fac["assets_value"],
        "annual_revenue": fac["annual_revenue"],
        "ebitda":        fac["ebitda"],
        "current_emissions_scope1": fac["current_emissions_scope1"],
        "current_emissions_scope2": fac["current_emissions_scope2"],
        "current_emissions_scope3": fac["current_emissions_scope3"],
    }
    for col_idx, col_name in enumerate(fac_cols, start=1):
        cell = ws2.cell(row=row_idx, column=col_idx, value=row_values[col_name])
        if col_name in _numeric_fac_cols:
            cell.number_format = "#,##0"

_autofit(ws2)


# ── Sheet 3: Hazard_Results ───────────────────────────────────────────────────
ws3 = wb.create_sheet("Hazard_Results")

hazard_cols = [
    "facility_id", "facility_name", "hazard_type", "mode",
    "risk_level", "probability", "potential_loss",
    "return_period_years", "climate_change_multiplier",
    "data_source", "has_api_warning", "cache_hit", "diff_pct",
]

_header_row(ws3, 1, hazard_cols)

# Build lookup: facility_id × hazard_type → static potential_loss
_static_loss_lookup: dict[tuple[str, str], float] = {}
for fac_r in r_static["facilities"]:
    for h in fac_r["hazards"]:
        _static_loss_lookup[(fac_r["facility_id"], h["hazard_type"])] = h.get(
            "potential_loss", 0
        )

# Build rows for both modes, then sort
hazard_rows: list[dict] = []

for mode_label, result in [("API", r_api), ("Static", r_static)]:
    for fac_r in result["facilities"]:
        fac_id   = fac_r["facility_id"]
        fac_name = fac_r["facility_name"]
        api_warn_hazards: set[str] = set()
        for w in fac_r.get("api_warnings", []):
            # Warning format: "hazard_type: API unavailable, ..."
            api_warn_hazards.add(w.split(":")[0].strip())

        for h in fac_r["hazards"]:
            hazard_type = h["hazard_type"]
            potential_loss = h.get("potential_loss", 0)

            # diff_pct only for API rows
            if mode_label == "API":
                static_loss = _static_loss_lookup.get((fac_id, hazard_type), 0)
                if static_loss > 0:
                    diff_pct: object = round(
                        (potential_loss - static_loss) / static_loss * 100, 1
                    )
                else:
                    diff_pct = "N/A"
            else:
                diff_pct = ""

            hazard_rows.append({
                "facility_id":              fac_id,
                "facility_name":            fac_name,
                "hazard_type":              hazard_type,
                "mode":                     mode_label,
                "risk_level":               h.get("risk_level", ""),
                "probability":              h.get("probability", ""),
                "potential_loss":           potential_loss,
                "return_period_years":      h.get("return_period_years", ""),
                "climate_change_multiplier": h.get("climate_change_multiplier", ""),
                "data_source":              h.get("data_source", ""),
                "has_api_warning":          hazard_type in api_warn_hazards,
                "cache_hit":                "N/A",
                "diff_pct":                 diff_pct,
            })

# Sort: facility_name, hazard_type, mode
hazard_rows.sort(key=lambda r: (r["facility_name"], r["hazard_type"], r["mode"]))

_risk_fill = {"High": _FILL_HIGH, "Medium": _FILL_MEDIUM, "Low": _FILL_LOW}
_risk_col_idx = hazard_cols.index("risk_level") + 1
_loss_col_idx = hazard_cols.index("potential_loss") + 1

for row_idx, row in enumerate(hazard_rows, start=2):
    for col_idx, col_name in enumerate(hazard_cols, start=1):
        cell = ws3.cell(row=row_idx, column=col_idx, value=row[col_name])
        if col_name == "potential_loss" and isinstance(row[col_name], (int, float)):
            cell.number_format = "#,##0"
        if col_name == "risk_level" and row[col_name] in _risk_fill:
            cell.fill = _risk_fill[row[col_name]]

_autofit(ws3)


# ── Sheet 4: Data_Source_Explanation ─────────────────────────────────────────
ws4 = wb.create_sheet("Data_Source_Explanation")

# Section 1
ws4.cell(row=1, column=1, value="데이터 소스별 hazard").font = _FONT_BOLD
_header_row(ws4, 2, ["hazard_type", "data_source", "설명", "한계"])

section1_rows = [
    (
        "flood",
        "open_meteo_era5",
        "자산 좌표로 ERA5 1994-2023 일강수량 요청 → 연최대치 추출 → Gumbel Type I MoM 피팅 → 재현기간별 침수심 → USACE 피해율 적용",
        "지형·배수 미반영. 강수→침수 변환이 단순 유출계수 방식. 상류 유역 기여 없음.",
    ),
    (
        "heatwave",
        "open_meteo_era5",
        "자산 좌표로 ERA5 일최고기온 요청 → 33°C 초과일 카운트 (KMA 기준)",
        "연속일수 기준 미적용(KMA는 2일+ 연속 요구). ERA5 0.1도 격자 평균값 사용.",
    ),
    (
        "drought",
        "open_meteo_era5",
        "자산 좌표로 ERA5 일강수량 요청 → 연간 최장 연속 무강수일(1mm 미만) 계산",
        "K-water 공식 가뭄 지수와 다를 수 있음. 지하수·저수지 반영 없음.",
    ),
    (
        "typhoon",
        "static_config",
        "KMA NTC 1951-2023 통계 기반 6개 권역 평균 상륙 빈도 사용",
        "자산 좌표 미반영. 200km 반경 기준. IBTrACS 미연동.",
    ),
    (
        "sea_level_rise",
        "static_config",
        "IPCC AR6 WG1 Ch.9 해수면 상승률 + 해안/내륙 이진 분류",
        "고도 데이터 미사용. 해안선 거리 미계산.",
    ),
]

for row_idx, vals in enumerate(section1_rows, start=3):
    for col_idx, val in enumerate(vals, start=1):
        cell = ws4.cell(row=row_idx, column=col_idx, value=val)
        cell.alignment = Alignment(wrap_text=True)

# Section 2
section2_start = 3 + len(section1_rows) + 2
ws4.cell(row=section2_start, column=1, value="API 실패 시 동작").font = _FONT_BOLD
_header_row(ws4, section2_start + 1, ["상황", "동작"])
ws4.cell(row=section2_start + 2, column=1, value="HTTP 429 또는 API 오류")
ws4.cell(
    row=section2_start + 2,
    column=2,
    value=(
        "Open-Meteo API 호출이 HTTP 429(속도 제한) 또는 기타 오류를 반환하면, "
        "모델은 오류를 발생시키지 않고 해당 hazard에 대해 static_config로 자동 폴백합니다. "
        "폴백 발생 시 해당 hazard 이름이 시설별 api_warnings 목록에 기록되고, "
        "최상위 결과의 api_warnings에도 집계됩니다. "
        "data_source 필드가 'open_meteo_era5' 대신 'static_config'로 표시됩니다."
    ),
)
ws4.cell(row=section2_start + 2, column=2).alignment = Alignment(wrap_text=True)

_autofit(ws4)
ws4.column_dimensions["C"].width = 55
ws4.column_dimensions["D"].width = 55
ws4.column_dimensions["B"].width = 55


# ── Sheet 5: Methodology_Notes ────────────────────────────────────────────────
ws5 = wb.create_sheet("Methodology_Notes")

method_cols = ["항목", "정의", "단위", "해석 시 주의사항"]
_header_row(ws5, 1, method_cols)

method_rows = [
    (
        "potential_loss",
        "각 재현기간별 예상손실을 발생확률로 가중합산한 연간 기대손실 (EAL). 영업중단 비용 포함",
        "USD",
        "단일 이벤트 손실이 아닌 연간 평균값",
    ),
    (
        "probability",
        (
            "해당 hazard의 연간 발생 확률. hazard별 정의 다름 "
            "(홍수: 최빈 재현기간 기준, 태풍: 연간 상륙 빈도, 폭염/가뭄: 일수 비율)"
        ),
        "0-1",
        "hazard 간 직접 비교 주의",
    ),
    (
        "return_period_years",
        "해당 강도의 이벤트가 평균적으로 몇 년에 한 번 발생하는지. 기후변화 배율 적용 후 값",
        "년",
        "100년 빈도 = 매년 1% 확률",
    ),
    (
        "climate_change_multiplier",
        "기후변화로 인한 빈도 또는 강도 증가 배율. IPCC AR6 WG1 Ch.11 기반",
        "배율",
        "홍수는 freq×intensity 복합, 태풍은 freq만 반영",
    ),
    (
        "data_source",
        "해당 hazard 계산에 사용된 데이터 출처",
        "-",
        "open_meteo_era5: 좌표 기반 실측. static_config: 권역 평균 추정치",
    ),
    (
        "static_config vs open_meteo_era5",
        (
            "static_config는 한국을 6개 권역으로 분류한 평균값 사용. "
            "open_meteo_era5는 자산 좌표의 ERA5 30년 데이터로 직접 계산"
        ),
        "-",
        "동일 권역 내 자산은 static_config에서 동일한 기준값을 가짐",
    ),
    (
        "GRESB 주의사항",
        (
            "GRESB 2025 Physical Risk 평가 시: "
            "(1) data_source가 open_meteo_era5인 hazard는 ERA5 근거 제시 가능. "
            "(2) 태풍·해수면 상승은 정적 권역 모델임을 방법론 섹션에 명시 권장. "
            "(3) potential_loss는 스크리닝 수준 추정값으로 site-specific 엔지니어링 평가 대체 불가"
        ),
        "-",
        "-",
    ),
]

for row_idx, vals in enumerate(method_rows, start=2):
    for col_idx, val in enumerate(vals, start=1):
        cell = ws5.cell(row=row_idx, column=col_idx, value=val)
        cell.alignment = Alignment(wrap_text=True)

_autofit(ws5)
for letter in ["A", "B", "C", "D"]:
    ws5.column_dimensions[letter].width = max(
        ws5.column_dimensions[letter].width, 40
    )


# ── Save ──────────────────────────────────────────────────────────────────────
output_dir = Path(__file__).resolve().parent.parent / "outputs"
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / "client_re_fund_physical_risk_feature_results.xlsx"

wb.save(output_path)

print(f"\n저장 완료: {output_path}")
print("시트 목록:")
for sheet_name in wb.sheetnames:
    print(f"  - {sheet_name}")
